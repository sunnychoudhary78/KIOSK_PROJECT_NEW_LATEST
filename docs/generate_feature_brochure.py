"""Build the Smart Kiosk desktop feature brochure (docs/Smart-Kiosk-Features.docx)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("Smart-Kiosk-Features.docx")

TEAL = "0F6A5A"
TEAL_BRIGHT = "2EC4A0"
GOLD = "E8C872"
INK = "1A2430"
MUTED = "5C6B76"
WHITE = "FFFFFF"
CARD = "F7FBFA"
MINT = "E7F4F0"

INK_RGB = RGBColor(0x1A, 0x24, 0x30)
MUTED_RGB = RGBColor(0x5C, 0x6B, 0x76)
TEAL_RGB = RGBColor(0x0F, 0x6A, 0x5A)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)
GOLD_RGB = RGBColor(0xE8, 0xC8, 0x72)
BRIGHT_RGB = RGBColor(0x2E, 0xC4, 0xA0)

# A4 content width at 1.7 cm side margins: 21.0 - 3.4 = 17.6 cm
CONTENT_CM = 17.6


def _hex(value: str) -> str:
    return value.replace("#", "").upper()


def set_run_font(run, name: str, size_pt: float, color: RGBColor, bold: bool = False) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.font.italic = False


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:shd"):
            tc_pr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(fill))
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:tcMar"):
            tc_pr.remove(child)
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)


def set_cell_borders(cell, **edges) -> None:
    """edges: top/left/bottom/right = (sz, color) or None to omit."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:tcBorders"):
            tc_pr.remove(child)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec is None:
            node.set(qn("w:val"), "nil")
        else:
            sz, color = spec
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), str(sz))
            node.set(qn("w:space"), "0")
            node.set(qn("w:color"), _hex(color))
        borders.append(node)
    tc_pr.append(borders)


def set_cell_width(cell, cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(tc_pr):
        if child.tag == qn("w:tcW"):
            tc_pr.remove(child)
    tc_w = OxmlElement("w:tcW")
    tc_w.set(qn("w:w"), str(int(cm * 567)))
    tc_w.set(qn("w:type"), "dxa")
    tc_pr.append(tc_w)


def set_table_widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(widths_cm)
    tbl_w.set(qn("w:w"), str(int(total * 567)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 567)))
        grid.append(col)
    tbl_pr.addnext(grid)

    for row in table.rows:
        tr = row._tr
        tr_pr = tr.get_or_add_trPr()
        cant = tr_pr.find(qn("w:cantSplit"))
        if cant is None:
            cant = OxmlElement("w:cantSplit")
            tr_pr.append(cant)
        for cell, width in zip(row.cells, widths_cm):
            set_cell_width(cell, width)


def clear_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)
    tbl_pr.append(borders)


def paragraph_spacing(p, before=0, after=0, line=None, line_rule="auto") -> None:
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE if line_rule == "auto" else WD_LINE_SPACING.EXACTLY
    if line is not None:
        pf.line_spacing = line


def add_text(p, text, size, color, bold=False, font="Calibri") -> None:
    run = p.add_run(text)
    set_run_font(run, font, size, color, bold)
    return run


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    set_run_font(run, "Calibri", 9, MUTED_RGB)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


SERVICES = [
    (
        "01",
        "Print",
        "Scan a QR code and upload from any phone browser — no app required. Or enter the code from the Smart Kiosk app. Preview the pages, then collect them from the printer.",
    ),
    (
        "02",
        "DigiLocker Print",
        "Sign in to DigiLocker on the kiosk, choose the government documents you need, preview them, and print.",
    ),
    (
        "03",
        "Well Being",
        "Check blood oxygen and heart rate with a fingertip reading, or take a body temperature. Results show on the screen in a few moments.",
    ),
    (
        "04",
        "Phone Charging",
        "Watch a short sponsored message, then plug in. The screen follows the charging session until you finish and return to the menu.",
    ),
    (
        "05",
        "Astrology",
        "See a birth chart in North or South Indian style, photograph a palm for insights, match two charts, or open today’s panchang.",
    ),
    (
        "06",
        "Fuel stations",
        "Find petrol, diesel, CNG, and EV stations nearby. Filter by fuel, open a map, and scan a QR code to continue directions on your phone.",
    ),
]

PUBLIC = [
    ("Home", "Tap Home when you are done. The next person starts with a clear screen."),
    ("Still there?", "If the kiosk sits idle, a short countdown returns it to the menu."),
    ("At a glance", "The home screen shows the time and local weather while you choose a service."),
    ("Out of paper", "If the printer runs out of paper, printing pauses with a plain message."),
]


def fill_card(cell, number: str, title: str, body: str) -> None:
    shade_cell(cell, CARD)
    set_cell_margins(cell, top=140, bottom=150, left=160, right=160)
    set_cell_borders(
        cell,
        top=(12, GOLD),
        left=(20, TEAL),
        bottom=(4, "D7E6E1"),
        right=(4, "D7E6E1"),
    )
    tc_pr = cell._tc.get_or_add_tcPr()
    v_align = tc_pr.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        tc_pr.append(v_align)
    v_align.set(qn("w:val"), "top")

    p_num = cell.paragraphs[0]
    p_num.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph_spacing(p_num, 0, 1, 1.0)
    add_text(p_num, number, 11, GOLD_RGB, bold=True)

    p_title = cell.add_paragraph()
    paragraph_spacing(p_title, 0, 4, 1.0)
    add_text(p_title, title, 16, TEAL_RGB, bold=True)

    p_body = cell.add_paragraph()
    paragraph_spacing(p_body, 0, 0, 1.12)
    add_text(p_body, body, 11, INK_RGB)


def set_exact_row_height(row, twips: int) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for child in list(tr_pr):
        if child.tag == qn("w:trHeight"):
            tr_pr.remove(child)
    height = OxmlElement("w:trHeight")
    height.set(qn("w:val"), str(twips))
    height.set(qn("w:hRule"), "exact")
    tr_pr.append(height)


def fill_public(cell, title: str, body: str) -> None:
    shade_cell(cell, MINT)
    set_cell_margins(cell, top=110, bottom=120, left=140, right=140)
    set_cell_borders(
        cell,
        top=(12, GOLD),
        left=(16, TEAL),
        bottom=(4, "D7E6E1"),
        right=(4, "D7E6E1"),
    )
    p_title = cell.paragraphs[0]
    paragraph_spacing(p_title, 0, 2, 1.0)
    add_text(p_title, title, 12, TEAL_RGB, bold=True)
    p_body = cell.add_paragraph()
    paragraph_spacing(p_body, 0, 0, 1.08)
    add_text(p_body, body, 10.5, INK_RGB)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK_RGB
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    core = doc.core_properties
    core.title = "Smart Kiosk — Desktop services"
    core.subject = "Feature overview of the Smart Kiosk desktop app"
    core.category = "Product brochure"
    core.comments = "Plain-language overview of visitor services on the Windows kiosk."

    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)

    # First page has no running header. Later pages carry the product name.
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph_spacing(hp, 0, 0, 1.0)
    add_text(hp, "Smart Kiosk", 10, TEAL_RGB, bold=True)
    add_text(hp, "   ·   Desktop services", 10, MUTED_RGB)
    # Gold underline via paragraph border
    p_pr = hp._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), GOLD)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""

    def paint_footer(footer) -> None:
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph_spacing(fp, 0, 0, 1.0)
        p_pr_f = fp._p.get_or_add_pPr()
        p_bdr_f = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "6")
        top.set(qn("w:space"), "8")
        top.set(qn("w:color"), TEAL)
        p_bdr_f.append(top)
        p_pr_f.append(p_bdr_f)
        fp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_CM), WD_TAB_ALIGNMENT.RIGHT)
        add_text(fp, "Smart Kiosk", 9, TEAL_RGB, bold=True)
        add_text(fp, "    Services in one touch", 9, MUTED_RGB)
        add_text(fp, "\t", 9, MUTED_RGB)
        add_page_field(fp)

    paint_footer(section.footer)
    paint_footer(section.first_page_footer)

    # Cover band
    cover = doc.add_table(rows=1, cols=1)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_table_borders(cover)
    set_table_widths(cover, [CONTENT_CM])
    band = cover.cell(0, 0)
    shade_cell(band, TEAL)
    set_cell_margins(band, top=220, bottom=200, left=220, right=220)
    set_cell_borders(band, top=None, left=None, bottom=(16, GOLD), right=None)

    kicker = band.paragraphs[0]
    paragraph_spacing(kicker, 0, 6, 1.0)
    add_text(kicker, "DESKTOP KIOSK", 11, GOLD_RGB, bold=True)

    title = band.add_paragraph()
    paragraph_spacing(title, 0, 2, 0.95)
    add_text(title, "Smart Kiosk", 36, WHITE_RGB, bold=True)

    promise = band.add_paragraph()
    paragraph_spacing(promise, 2, 8, 1.0)
    add_text(promise, "Services in one touch", 16, BRIGHT_RGB, bold=True)

    intro_in = band.add_paragraph()
    paragraph_spacing(intro_in, 4, 0, 1.15)
    add_text(
        intro_in,
        "A public touchscreen for everyday tasks. Printing, health checks, phone charging, astrology, and nearby fuel — each one a short visit, with no training required.",
        12,
        WHITE_RGB,
    )

    spacer = doc.add_paragraph()
    paragraph_spacing(spacer, 8, 2, 1.0)
    add_text(spacer, "WHAT YOU CAN DO", 11, TEAL_RGB, bold=True)

    lead = doc.add_paragraph()
    paragraph_spacing(lead, 0, 8, 1.12)
    add_text(
        lead,
        "Seven services sit on the home screen. Touch a tile to begin. When you are finished, Home brings the kiosk back for the next person.",
        12,
        INK_RGB,
    )

    gap = 0.35
    card_w = (CONTENT_CM - gap) / 2
    widths = [card_w, gap, card_w]

    grid = doc.add_table(rows=5, cols=3)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_table_borders(grid)
    set_table_widths(grid, widths)

    pairs = [(0, SERVICES[0], SERVICES[1]), (2, SERVICES[2], SERVICES[3]), (4, SERVICES[4], SERVICES[5])]
    for row_idx, left, right in pairs:
        fill_card(grid.cell(row_idx, 0), *left)
        fill_card(grid.cell(row_idx, 2), *right)
        mid = grid.cell(row_idx, 1)
        shade_cell(mid, WHITE)
        set_cell_borders(mid)
        set_cell_margins(mid, 0, 0, 0, 0)
        mid.paragraphs[0].text = ""

    for spacer_row in (1, 3):
        set_exact_row_height(grid.rows[spacer_row], 160)
        for col in range(3):
            cell = grid.cell(spacer_row, col)
            shade_cell(cell, WHITE)
            set_cell_borders(cell)
            set_cell_margins(cell, 0, 0, 0, 0)
            cell.paragraphs[0].text = ""
            paragraph_spacing(cell.paragraphs[0], 0, 0, 1.0)

    bridge = doc.add_paragraph()
    paragraph_spacing(bridge, 14, 4, 1.0)
    bridge.paragraph_format.keep_with_next = True
    add_text(bridge, "BUILT FOR PUBLIC USE", 11, TEAL_RGB, bold=True)

    note = doc.add_paragraph()
    paragraph_spacing(note, 0, 8, 1.12)
    note.paragraph_format.keep_with_next = True
    add_text(
        note,
        "The same screen stays simple for every visitor who walks up to it.",
        12,
        MUTED_RGB,
    )

    pub_gap = 0.3
    pub_w = (CONTENT_CM - pub_gap) / 2
    pub_widths = [pub_w, pub_gap, pub_w]
    public = doc.add_table(rows=3, cols=3)
    public.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_table_borders(public)
    set_table_widths(public, pub_widths)

    fill_public(public.cell(0, 0), *PUBLIC[0])
    fill_public(public.cell(0, 2), *PUBLIC[1])
    fill_public(public.cell(2, 0), *PUBLIC[2])
    fill_public(public.cell(2, 2), *PUBLIC[3])
    set_exact_row_height(public.rows[1], 140)
    for row_idx, col in ((0, 1), (2, 1), (1, 0), (1, 1), (1, 2)):
        cell = public.cell(row_idx, col)
        shade_cell(cell, WHITE)
        set_cell_borders(cell)
        set_cell_margins(cell, 0, 0, 0, 0)
        cell.paragraphs[0].text = ""

    local_kicker = doc.add_paragraph()
    local_kicker.paragraph_format.page_break_before = True
    local_kicker.paragraph_format.keep_with_next = True
    paragraph_spacing(local_kicker, 0, 2, 1.0)
    add_text(local_kicker, "07  ·  LOCAL SERVICES", 11, TEAL_RGB, bold=True)

    local_lead = doc.add_paragraph()
    local_lead.paragraph_format.keep_with_next = True
    paragraph_spacing(local_lead, 0, 8, 1.12)
    add_text(
        local_lead,
        "Find a person for an everyday job. Each listing shows the name, phone number, and location.",
        12,
        INK_RGB,
    )

    trades = [
        "AC, RO and washing machines",
        "Two-wheeler mechanics",
        "Carpenters",
        "Electricians",
        "Painters",
        "Tailors",
        "Scrap dealers",
        "Gas-stove repair",
        "Pigeon netting",
        "Puncture repair",
    ]
    trade_gap = 0.35
    trade_w = (CONTENT_CM - trade_gap) / 2
    trade_table = doc.add_table(rows=9, cols=3)
    trade_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    clear_table_borders(trade_table)
    set_table_widths(trade_table, [trade_w, trade_gap, trade_w])
    for index, label in enumerate(trades):
        row = (index // 2) * 2
        col = 0 if index % 2 == 0 else 2
        cell = trade_table.cell(row, col)
        shade_cell(cell, CARD)
        set_cell_margins(cell, top=90, bottom=90, left=140, right=120)
        set_cell_borders(
            cell,
            top=(8, GOLD),
            left=(16, TEAL),
            bottom=(4, "D7E6E1"),
            right=(4, "D7E6E1"),
        )
        paragraph_spacing(cell.paragraphs[0], 0, 0, 1.0)
        add_text(cell.paragraphs[0], label, 13, TEAL_RGB, bold=True)
    for spacer_row in (1, 3, 5, 7):
        set_exact_row_height(trade_table.rows[spacer_row], 120)
    for row_idx in range(9):
        for col in range(3):
            if row_idx % 2 == 1 or col == 1:
                cell = trade_table.cell(row_idx, col)
                shade_cell(cell, WHITE)
                set_cell_borders(cell)
                set_cell_margins(cell, 0, 0, 0, 0)
                cell.paragraphs[0].text = ""

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph_spacing(closing, 14, 0, 1.1)
    add_text(closing, "Smart Kiosk", 11, TEAL_RGB, bold=True)
    add_text(closing, "  ·  Windows desktop kiosk", 11, MUTED_RGB)

    # Drop the empty default paragraph Word inserts before the first table if it is blank
    body = doc.element.body
    first_p = body.find(qn("w:p"))
    if first_p is not None and not "".join(first_p.itertext()).strip():
        nxt = first_p.getnext()
        if nxt is not None and nxt.tag == qn("w:tbl"):
            body.remove(first_p)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
